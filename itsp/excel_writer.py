"""openpyxl ile incident_report.xlsx üretimi."""

from __future__ import annotations

import logging
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as cfg
from .scraper import Incident

logger = logging.getLogger(__name__)

HEADERS = [
    "Incident Number",
    "Status",
    "Priority",
    "Assignment Group",
    "Last Comment Date",
    "Last Comment Author",
    "Last Comment Side",
    "Working Days Since",
    "Reminder Stage",
    "Action",
]

_COLUMN_WIDTHS = [18, 22, 12, 24, 20, 22, 16, 16, 26, 40]


def write_report(incidents: Iterable[Incident], path=None) -> str:
    """Incident listesini Excel'e yazar ve dosya yolunu döndürür."""
    out_path = str(path or cfg.REPORT_FILE)

    wb = Workbook()
    ws = wb.active
    ws.title = "Incidents"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = _COLUMN_WIDTHS[col - 1]
    ws.freeze_panes = "A2"

    row = 2
    for inc in incidents:
        date_str = (
            inc.last_comment_date.strftime("%Y-%m-%d %H:%M")
            if inc.last_comment_date
            else inc.last_comment_date_raw
        )
        values = [
            inc.number,
            inc.status,
            inc.priority,
            inc.assignment_group,
            date_str,
            inc.last_comment_author,
            inc.side,
            inc.working_days_since if inc.working_days_since is not None else "",
            inc.reminder_stage,
            inc.action,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    wb.save(out_path)
    logger.info("Rapor yazıldı: %s (%d kayıt)", out_path, row - 2)
    return out_path
